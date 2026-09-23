"""Plantilla descargable y carga masiva del directorio de Empleados vía Excel.

Generado desde cero con openpyxl (no rellena ninguna plantilla institucional
existente, porque `Empleado` es un modelo propio de esta aplicación, no un formato
FO-860 del SGSI) — a diferencia de los generadores de auditoría, que sí rellenan
formatos reales del usuario.
"""

import io

import openpyxl
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from .models import Empleado

COLUMNAS = ['Nombre completo', 'Cargo', 'Correo', 'Activo (Sí/No)']
ANCHOS_COLUMNA = [32, 28, 34, 16]
TEAL_OSCURO = 'FF0B5C53'
FILAS_CON_VALIDACION = 500


def generar_plantilla_empleados_xlsx():
    """Devuelve un BytesIO con una plantilla .xlsx en blanco, lista para diligenciar
    y volver a subir con "Importar empleados"."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Empleados'

    # El ejemplo va como comentario de celda, no como datos reales en la fila 2:
    # si estuviera como fila de datos, alguien que suba la plantilla sin borrarla
    # crearía un empleado ficticio "Juana Pérez Gómez" de verdad (se probó y pasó).
    comentarios = {
        1: 'Obligatorio. Ejemplo: Juana Pérez Gómez',
        2: 'Opcional. Ejemplo: Coordinadora de Calidad',
        3: 'Opcional. Ejemplo: juana.perez@epsianaswayuu.com',
        4: 'Opcional — si se deja vacío, se carga como Activo. Escribe Sí o No.',
    }
    for col, titulo in enumerate(COLUMNAS, start=1):
        celda = ws.cell(row=1, column=col, value=titulo)
        celda.font = Font(bold=True, color='FFFFFFFF')
        celda.fill = PatternFill('solid', fgColor=TEAL_OSCURO)
        celda.alignment = Alignment(vertical='center')
        celda.comment = Comment(comentarios[col], 'SGSI')

    validacion_activo = DataValidation(type='list', formula1='"Sí,No"', allow_blank=True)
    ws.add_data_validation(validacion_activo)
    validacion_activo.add(f'D2:D{FILAS_CON_VALIDACION}')

    for i, ancho in enumerate(ANCHOS_COLUMNA, start=1):
        ws.column_dimensions[get_column_letter(i)].width = ancho
    ws.freeze_panes = 'A2'

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def _texto(valor):
    return str(valor).strip() if valor not in (None, '') else ''


def _parsear_activo(valor):
    if valor in (None, ''):
        return True
    texto = str(valor).strip().lower()
    return texto not in ('no', 'false', '0', 'inactivo')


def importar_empleados_desde_xlsx(archivo):
    """Lee la plantilla diligenciada y crea/actualiza Empleado por fila.

    Empareja por correo (si la fila trae uno) y si no, por nombre completo exacto —
    así reimportar el mismo archivo actualiza en vez de duplicar. Devuelve un resumen
    con conteos y, si hubo filas sin nombre, el detalle de cuáles se omitieron."""
    wb = openpyxl.load_workbook(archivo, data_only=True)
    ws = wb.active

    creados = 0
    actualizados = 0
    omitidos = []

    for indice, fila in enumerate(ws.iter_rows(min_row=2, max_col=4, values_only=True), start=2):
        if not fila or all(v in (None, '') for v in fila):
            continue
        nombre = _texto(fila[0] if len(fila) > 0 else None)
        cargo = _texto(fila[1] if len(fila) > 1 else None)
        correo = _texto(fila[2] if len(fila) > 2 else None)
        activo = _parsear_activo(fila[3] if len(fila) > 3 else None)

        if not nombre:
            omitidos.append(f'Fila {indice}: sin nombre completo, se omitió.')
            continue

        empleado = Empleado.objects.filter(correo__iexact=correo).first() if correo else None
        if not empleado:
            empleado = Empleado.objects.filter(nombre_completo__iexact=nombre).first()

        if empleado:
            empleado.nombre_completo = nombre
            empleado.cargo = cargo
            empleado.correo = correo
            empleado.activo = activo
            empleado.save()
            actualizados += 1
        else:
            Empleado.objects.create(nombre_completo=nombre, cargo=cargo, correo=correo, activo=activo)
            creados += 1

    return {'creados': creados, 'actualizados': actualizados, 'omitidos': omitidos}
