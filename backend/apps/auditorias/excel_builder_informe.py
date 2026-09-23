"""Genera el Informe de Auditoría Interna (FO-860-22) como un .xlsx real, rellenando
la plantilla institucional oficial (excel_templates/FO-860-22_informe_auditoria_interna.xlsx
— copia exacta del formato que el usuario aportó) en vez de reconstruirla desde cero,
mismo criterio que ya se usó para el acta de Revisión por la Dirección y el informe en
Word (ver apps/revisiones/docx_builder.py) — pero aquí el usuario pidió explícitamente
el formato Excel real, no Word, para este documento en particular.

La plantilla en blanco trae pre-armados: 3 bloques de "PROCESO" (filas 17, 28, 42, con
un encabezado de página repetido en medio) y 11 filas de la tabla de consolidado
(56-66). Como una auditoría real casi siempre tiene más de 3 procesos (la de diciembre
2025 tiene 12), este generador **reconstruye la zona de bloques usando SIEMPRE el
primer bloque (filas 17-27) como plantilla de estilo**, clonándolo tantas veces como
haga falta — no se reutilizan los bloques 2 y 3 originales (su espaciado y alturas de
fila no son uniformes entre sí, probablemente ajustados a mano por quien armó el
Excel) ni el encabezado de "página 2" intermedio, para mantener un patrón predecible.

openpyxl NO ajusta las celdas combinadas (merged cells) al insertar o eliminar filas
— es una limitación documentada de la librería, no un descuido. Por eso todas las
fusiones se desarman al principio y se reconstruyen al final, sobre las filas ya en su
posición definitiva.
"""

import io
from copy import copy

import openpyxl
from django.conf import settings

from .models import ItemVerificacionAuditoria

RUTA_PLANTILLA = settings.BASE_DIR / 'apps' / 'auditorias' / 'excel_templates' / 'FO-860-22_informe_auditoria_interna.xlsx'

FILA_BLOQUE_PLANTILLA = 17
ALTO_BLOQUE = 11  # filas 17..27 inclusive
FILA_CONSOLIDADO_LABEL_ORIGINAL = 53
FILA_CONSOLIDADO_HEADER_ORIGINAL = 55
FILA_CONSOLIDADO_DATOS_INICIO_ORIGINAL = 56
FILAS_CONSOLIDADO_PREARMADAS = 11  # filas 56..66
COLUMNAS_A_W = range(1, 24)


def _copiar_estilo_fila(ws, fila_origen, fila_destino, copiar_valor=False):
    if fila_origen in ws.row_dimensions:
        ws.row_dimensions[fila_destino].height = ws.row_dimensions[fila_origen].height
    for col in COLUMNAS_A_W:
        co = ws.cell(row=fila_origen, column=col)
        cd = ws.cell(row=fila_destino, column=col)
        cd.font = copy(co.font)
        cd.fill = copy(co.fill)
        cd.border = copy(co.border)
        cd.alignment = copy(co.alignment)
        cd.number_format = co.number_format
        if copiar_valor:
            cd.value = co.value


def _extraer_y_desarmar_merges(ws):
    """Devuelve la lista de fusiones (min_col, min_fila, max_col, max_fila) y las
    quita todas del sheet — necesario antes de insertar/eliminar filas."""
    rangos = [rng.bounds for rng in list(ws.merged_cells.ranges)]
    for rng in list(ws.merged_cells.ranges):
        ws.unmerge_cells(str(rng))
    return rangos


def _rearmar_merges(ws, rangos, mapa_filas):
    """mapa_filas: función fila_original -> fila_nueva (o None si la fila desapareció).
    Vuelve a fusionar cada rango cuya fila inicial y final sigan existiendo."""
    for min_col, min_fila, max_col, max_fila in rangos:
        nueva_min = mapa_filas(min_fila)
        nueva_max = mapa_filas(max_fila)
        if nueva_min is None or nueva_max is None:
            continue
        ws.merge_cells(start_row=nueva_min, start_column=min_col, end_row=nueva_max, end_column=max_col)


def _clonar_bloque_proceso(ws, fila_destino_inicio):
    """Copia estilo + etiquetas fijas del bloque plantilla (17-27) a un bloque de 11
    filas en blanco ya presente en fila_destino_inicio."""
    for i in range(ALTO_BLOQUE):
        fila_o = FILA_BLOQUE_PLANTILLA + i
        fila_d = fila_destino_inicio + i
        es_fila_con_etiqueta = i in (0, 2, 5, 8)  # PROCESO/RESPONSABLE, FORTALEZAS, ASPECTOS, NO CONFORMIDADES
        _copiar_estilo_fila(ws, fila_o, fila_d, copiar_valor=es_fila_con_etiqueta)


def generar_informe_auditoria_xlsx(auditoria):
    """Devuelve un BytesIO con el .xlsx del informe, listo para servir como descarga."""
    wb = openpyxl.load_workbook(str(RUTA_PLANTILLA))
    ws = wb.active

    rangos_merge, _ = _extraer_y_desarmar_merges(ws), None

    # --- 1) Encabezado: objetivo/alcance/criterios/equipo/fechas ---
    ws['I5'] = auditoria.objetivo or ''
    ws['I7'] = auditoria.alcance or ''
    ws['I9'] = auditoria.criterios or ''
    nombres_equipo = [auditoria.auditor_lider.nombre_completo] if auditoria.auditor_lider_id else []
    nombres_equipo += [e.nombre_completo for e in auditoria.equipo_auditor.all()]
    ws['I11'] = ', '.join(dict.fromkeys(nombres_equipo)) or '—'
    ws['I13'] = str(auditoria.fecha_auditoria) if auditoria.fecha_auditoria else '—'
    ws['S13'] = str(auditoria.fecha_elaboracion_informe) if auditoria.fecha_elaboracion_informe else '—'

    # --- 2) Agrupar hallazgos por proceso/tema, igual criterio que el Word ---
    grupos = []
    vistos = set()
    for hallazgo in auditoria.hallazgos.prefetch_related('procesos').order_by('id'):
        proceso = hallazgo.procesos.first()
        etiqueta = proceso.nombre if proceso else (hallazgo.tema or None)
        if not etiqueta or etiqueta in vistos:
            continue
        vistos.add(etiqueta)
        grupos.append((etiqueta, proceso))

    def responsable_de(proceso, etiqueta):
        filtro = {'proceso': proceso} if proceso else {'tema': etiqueta}
        sesion = auditoria.sesiones.filter(**filtro).exclude(auditado='').first()
        return sesion.auditado if sesion else '—'

    def hallazgos_de(proceso, etiqueta, codigo_tipo):
        filtro = {'procesos': proceso} if proceso else {'tema': etiqueta}
        return list(
            auditoria.hallazgos.filter(tipos__codigo=codigo_tipo, **filtro)
            .values_list('descripcion', flat=True).distinct()
        )

    # --- 3) Eliminar los bloques 2 y 3 originales + el encabezado de "página 2"
    #         intermedio (filas 28 a 52), dejando el bloque 1 (17-27) intacto y todo
    #         lo de consolidado en adelante (antes fila 53) pegado justo después.
    ws.delete_rows(28, 52 - 28 + 1)

    n_grupos = len(grupos)
    filas_extra_bloques = max(0, n_grupos - 1) * ALTO_BLOQUE
    if filas_extra_bloques:
        ws.insert_rows(28, filas_extra_bloques)
        for i in range(n_grupos - 1):
            _clonar_bloque_proceso(ws, 28 + i * ALTO_BLOQUE)

    # Posiciones ya definitivas de la zona de bloques (block1 sigue en 17; los N-1
    # clones ocupan 28 en adelante).
    filas_bloques = [FILA_BLOQUE_PLANTILLA + i * ALTO_BLOQUE for i in range(n_grupos)]

    for (etiqueta, proceso), fila in zip(grupos, filas_bloques):
        ws.cell(row=fila, column=4).value = etiqueta  # D
        ws.cell(row=fila, column=21).value = responsable_de(proceso, etiqueta)  # U
        fortalezas = hallazgos_de(proceso, etiqueta, 'FORT')
        aspectos = hallazgos_de(proceso, etiqueta, 'AM')
        no_conformidades = hallazgos_de(proceso, etiqueta, 'NC')
        ws.cell(row=fila + 3, column=2).value = '\n'.join(f'* {t}' for t in fortalezas) or 'N.D.'
        ws.cell(row=fila + 6, column=2).value = '\n'.join(f'* {t}' for t in aspectos) or 'N.D.'
        ws.cell(row=fila + 9, column=2).value = '\n'.join(f'* {t}' for t in no_conformidades) or 'N.D.'

    # --- 4) Consolidado: la plantilla trae 11 filas de datos; si hay más grupos,
    #         clonar filas adicionales usando la primera fila de datos como estilo.
    desplazamiento_bloques = filas_extra_bloques  # todo lo de abajo se corrió esto
    fila_consolidado_datos_inicio = FILA_CONSOLIDADO_DATOS_INICIO_ORIGINAL - (52 - 28 + 1) + desplazamiento_bloques

    filas_extra_consolidado = max(0, n_grupos - FILAS_CONSOLIDADO_PREARMADAS)
    if filas_extra_consolidado:
        fila_insercion = fila_consolidado_datos_inicio + FILAS_CONSOLIDADO_PREARMADAS
        ws.insert_rows(fila_insercion, filas_extra_consolidado)
        for i in range(filas_extra_consolidado):
            _copiar_estilo_fila(ws, fila_consolidado_datos_inicio, fila_insercion + i)

    for indice, (etiqueta, _proceso) in enumerate(grupos):
        fila = fila_consolidado_datos_inicio + indice
        f = len(hallazgos_de(grupos[indice][1], etiqueta, 'FORT'))
        a = len(hallazgos_de(grupos[indice][1], etiqueta, 'AM'))
        n = len(hallazgos_de(grupos[indice][1], etiqueta, 'NC'))
        ws.cell(row=fila, column=2).value = etiqueta  # B
        ws.cell(row=fila, column=11).value = f  # K
        ws.cell(row=fila, column=17).value = a  # Q
        ws.cell(row=fila, column=22).value = n  # V

    # --- 5) Reconstruir las fusiones: primero las que quedaron dentro del bloque 1
    #         (nunca se movió), luego todo lo que estaba después del área borrada
    #         (28..52), corrido por el neto de filas insertadas/eliminadas.
    neto_bloques = desplazamiento_bloques - (52 - 28 + 1)
    neto_consolidado = filas_extra_consolidado

    def mapa_filas(fila_original):
        if fila_original <= 27:
            return fila_original
        if 28 <= fila_original <= 52:
            return None  # zona eliminada (bloques 2/3 originales + encabezado p2)
        fila_tras_bloques = fila_original + neto_bloques
        if fila_tras_bloques >= fila_consolidado_datos_inicio + FILAS_CONSOLIDADO_PREARMADAS:
            return fila_tras_bloques + neto_consolidado
        return fila_tras_bloques

    _rearmar_merges(ws, rangos_merge, mapa_filas)

    # Fusiones nuevas para los bloques de proceso clonados (2º en adelante) y las
    # filas de consolidado clonadas — mismo patrón relativo que el bloque/fila 1.
    rangos_bloque1 = [
        (c0, f0, c1, f1) for c0, f0, c1, f1 in rangos_merge
        if FILA_BLOQUE_PLANTILLA <= f0 and f1 <= FILA_BLOQUE_PLANTILLA + ALTO_BLOQUE - 1
    ]
    for fila_destino in filas_bloques[1:]:
        offset = fila_destino - FILA_BLOQUE_PLANTILLA
        for c0, f0, c1, f1 in rangos_bloque1:
            ws.merge_cells(start_row=f0 + offset, start_column=c0, end_row=f1 + offset, end_column=c1)

    rango_fila_consolidado1 = [
        (c0, f0, c1, f1) for c0, f0, c1, f1 in rangos_merge
        if f0 == FILA_CONSOLIDADO_DATOS_INICIO_ORIGINAL and f1 == FILA_CONSOLIDADO_DATOS_INICIO_ORIGINAL
    ]
    for indice in range(filas_extra_consolidado):
        fila_destino = fila_consolidado_datos_inicio + FILAS_CONSOLIDADO_PREARMADAS + indice
        offset = fila_destino - FILA_CONSOLIDADO_DATOS_INICIO_ORIGINAL
        for c0, f0, c1, f1 in rango_fila_consolidado1:
            ws.merge_cells(start_row=f0 + offset, start_column=c0, end_row=f1 + offset, end_column=c1)

    # --- 6) Conclusiones y firmas (posiciones ya corridas por mapa_filas) ---
    fila_conclusiones_valor = mapa_filas(69)
    if fila_conclusiones_valor:
        ws.cell(row=fila_conclusiones_valor, column=2).value = auditoria.conclusiones_generales or 'Sin conclusiones registradas.'
    fila_nombre_auditor = mapa_filas(71)
    if fila_nombre_auditor and auditoria.auditor_lider_id:
        ws.cell(row=fila_nombre_auditor, column=11).value = auditoria.auditor_lider.nombre_completo
    fila_aprueba = mapa_filas(76)
    if fila_aprueba and auditoria.aprobado_por_id:
        ws.cell(row=fila_aprueba, column=4).value = auditoria.aprobado_por.nombre_completo
    fila_cargo = mapa_filas(78)
    if fila_cargo and auditoria.aprobado_por_id and auditoria.aprobado_por.cargo:
        ws.cell(row=fila_cargo, column=4).value = auditoria.aprobado_por.cargo

    total_items = auditoria.items_verificacion.count()
    if total_items:
        conformidades = auditoria.items_verificacion.filter(
            tipo_hallazgo=ItemVerificacionAuditoria.TipoHallazgoChecklist.CONFORMIDAD
        ).count()
        fila_cobertura = (fila_conclusiones_valor or 0) - 1
        if fila_cobertura > 0:
            ws.cell(row=fila_cobertura, column=2).value = (
                f'Cobertura de la lista de verificación: {total_items} elementos revisados, '
                f'{conformidades} en conformidad.'
            )

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer
