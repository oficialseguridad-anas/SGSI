"""Genera el acta de Revisión por la Dirección como un .docx real, rellenando el
membrete institucional oficial de ANAS WAYUU EPSI (docx_templates/membrete_institucional.docx
— provisto por el usuario) en vez de reconstruirlo desde cero.

Se abandonó la generación en PDF (xhtml2pdf): al usar `@frame` para lograr una franja
lateral fija por página, xhtml2pdf dejaba de paginar automáticamente el contenido largo
(confirmado con pypdf: una acta que debía ocupar 2 páginas generaba 1 sola, con todo el
contenido superpuesto e ilegible). Word maneja encabezados/pies/logos repetidos por
página de forma nativa, sin ese problema — y como esta plantilla YA tiene el membrete
real armado (con la franja "Vigilado Supersalud" en su posición correcta), no hace falta
reconstruir nada de eso: solo se agrega contenido al cuerpo del documento, dejando el
encabezado/pie de la plantilla intactos.

No se usa `Document.add_heading()`: la plantilla no define estilos "Heading N" (se
confirmó revisando `doc.styles` — solo trae 'Normal', 'Header', 'Footer', 'Table Grid',
etc.), así que `add_heading()` fallaría con KeyError. Todos los títulos se arman con
párrafos + formato de texto manual (negrita/color/tamaño de fuente).
"""

import io

from django.conf import settings
from django.utils import timezone
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

RUTA_PLANTILLA = settings.BASE_DIR / 'apps' / 'revisiones' / 'docx_templates' / 'membrete_institucional.docx'

TEAL_OSCURO = RGBColor(0x0B, 0x5C, 0x53)
GRIS_SECUNDARIO = RGBColor(0x52, 0x51, 0x4E)
GRIS_CLARO = RGBColor(0x89, 0x87, 0x81)
BLANCO = RGBColor(0xFF, 0xFF, 0xFF)


def _sombrear_celda(celda, color_hex):
    """python-docx no tiene API de alto nivel para el color de fondo de una celda —
    hay que agregar el elemento `w:shd` directamente al XML de la celda."""
    propiedades = celda._tc.get_or_add_tcPr()
    sombreado = OxmlElement('w:shd')
    sombreado.set(qn('w:val'), 'clear')
    sombreado.set(qn('w:color'), 'auto')
    sombreado.set(qn('w:fill'), color_hex)
    propiedades.append(sombreado)


def _titulo_principal(doc, texto):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    run = p.add_run(texto)
    run.bold = True
    run.font.size = Pt(15)
    run.font.color.rgb = TEAL_OSCURO
    return p


def _subtitulo_principal(doc, texto):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(10)
    run = p.add_run(texto)
    run.font.size = Pt(10)
    run.font.color.rgb = GRIS_SECUNDARIO
    return p


def _titulo_gran_seccion(doc, texto):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(texto)
    run.bold = True
    run.font.size = Pt(13)
    run.font.color.rgb = TEAL_OSCURO
    return p


def _titulo_seccion(doc, texto):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(texto)
    run.bold = True
    run.font.size = Pt(11)
    run.font.color.rgb = TEAL_OSCURO
    return p


def _texto_o_vacio(doc, texto, mensaje_vacio='Sin observaciones registradas.'):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    if texto:
        run = p.add_run(texto)
        run.font.size = Pt(10.5)
    else:
        run = p.add_run(mensaje_vacio)
        run.italic = True
        run.font.size = Pt(10)
        run.font.color.rgb = GRIS_CLARO
    return p


def _cifras(doc, texto):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(texto)
    run.font.size = Pt(9)
    run.font.color.rgb = GRIS_SECUNDARIO
    return p


def _fila_info(tabla, etiqueta, valor):
    fila = tabla.add_row()
    celda_etiqueta, celda_valor = fila.cells
    celda_etiqueta.text = etiqueta
    for run in celda_etiqueta.paragraphs[0].runs:
        run.bold = True
        run.font.color.rgb = TEAL_OSCURO
    _sombrear_celda(celda_etiqueta, 'F4F7F6')
    celda_valor.text = valor


def generar_acta_docx(revision):
    """Devuelve un BytesIO con el .docx del acta, listo para servir como descarga."""
    doc = Document(str(RUTA_PLANTILLA))
    resumen = revision.resumen_datos

    _titulo_principal(doc, f'ACTA DE REVISIÓN POR LA DIRECCIÓN — {revision.periodo}')
    _subtitulo_principal(
        doc, 'Sistema de Gestión de Seguridad de la Información — Cláusula 9.3 de ISO/IEC 27001:2022'
    )

    tabla_info = doc.add_table(rows=0, cols=2)
    tabla_info.style = 'Table Grid'
    tabla_info.columns[0].width = Pt(140)
    _fila_info(tabla_info, 'Fecha de revisión', str(revision.fecha_revision))
    _fila_info(tabla_info, 'Estado', 'Finalizada' if revision.finalizada else 'Borrador')
    preside_texto = revision.preside.nombre_completo
    if revision.preside.cargo:
        preside_texto += f' ({revision.preside.cargo})'
    _fila_info(tabla_info, 'Preside', preside_texto)
    asistentes = ', '.join(a.nombre_completo for a in revision.asistentes.all()) or '—'
    _fila_info(tabla_info, 'Asistentes', asistentes)
    _fila_info(tabla_info, 'Lugar / modalidad', revision.lugar_modalidad or '—')

    _titulo_gran_seccion(doc, 'Entradas de la revisión (cláusula 9.3.2)')

    _titulo_seccion(doc, 'a) Estado de las acciones de revisiones por la dirección previas')
    _texto_o_vacio(doc, revision.estado_acciones_previas)

    _titulo_seccion(doc, 'b) Cambios en las cuestiones externas e internas')
    _texto_o_vacio(doc, revision.cambios_cuestiones_externas_internas)

    _titulo_seccion(doc, 'c) Cambios en las necesidades y expectativas de las partes interesadas')
    _texto_o_vacio(doc, revision.cambios_partes_interesadas)

    _titulo_seccion(doc, 'd.1) No conformidades y acciones correctivas')
    _cifras(doc, (
        f"{resumen['hallazgos']['total']} hallazgos totales — "
        f"{resumen['hallazgos']['por_estado']['ABIERTA']} abiertos, "
        f"{resumen['hallazgos']['por_estado']['EN_PROCESO']} en proceso, "
        f"{resumen['hallazgos']['por_estado']['CERRADA']} cerrados."
    ))
    _texto_o_vacio(doc, revision.desempeno_no_conformidades)

    _titulo_seccion(doc, 'd.2) Resultados de seguimiento y medición (indicadores)')
    _cifras(doc, (
        f"{resumen['indicadores']['total']} indicadores — {resumen['indicadores']['al_dia']} al día, "
        f"{resumen['indicadores']['cumple']} cumplen su meta."
    ))
    _texto_o_vacio(doc, revision.desempeno_seguimiento_medicion)

    _titulo_seccion(doc, 'd.3) Resultados de auditoría')
    _texto_o_vacio(doc, revision.desempeno_auditorias)

    _titulo_seccion(doc, 'd.4) Cumplimiento de los objetivos de seguridad de la información')
    _cifras(doc, (
        f"{resumen['objetivos']['total']} objetivos — "
        f"{resumen['objetivos']['actividades_por_estado']['COMPLETADA']} actividades completadas, "
        f"{resumen['objetivos']['actividades_por_estado']['PENDIENTE']} pendientes, "
        f"{resumen['objetivos']['actividades_por_estado']['VENCIDA']} vencidas."
    ))
    _texto_o_vacio(doc, revision.desempeno_objetivos)

    _titulo_seccion(doc, 'e) Retroalimentación de las partes interesadas')
    _texto_o_vacio(doc, revision.retroalimentacion_partes_interesadas)

    _titulo_seccion(doc, 'f) Resultados de la valoración de riesgos y estado del plan de tratamiento')
    _cifras(doc, (
        f"{resumen['riesgos']['total']} riesgos — {resumen['riesgos']['por_nivel']['CRITICO']} críticos, "
        f"{resumen['riesgos']['por_nivel']['ALTO']} altos, {resumen['riesgos']['por_nivel']['MEDIO']} medios, "
        f"{resumen['riesgos']['por_nivel']['BAJO']} bajos. "
        f"{resumen['riesgos']['tratamientos_pendientes']} tratamientos pendientes, "
        f"{resumen['riesgos']['tratamientos_vencidos']} vencidos. "
        f"{resumen['incidentes']['total']} incidentes/eventos registrados."
    ))
    _texto_o_vacio(doc, revision.resultados_riesgos)

    _titulo_seccion(doc, 'g) Oportunidades de mejora continua')
    _texto_o_vacio(doc, revision.oportunidades_mejora)

    _titulo_gran_seccion(doc, 'Conclusiones generales')
    _texto_o_vacio(doc, revision.conclusiones_generales, mensaje_vacio='Sin conclusiones registradas.')

    _titulo_gran_seccion(doc, 'Salidas — compromisos y decisiones (cláusula 9.3.3)')
    compromisos = list(revision.compromisos.all())
    if compromisos:
        tabla_c = doc.add_table(rows=1, cols=4)
        tabla_c.style = 'Table Grid'
        encabezados = ['Compromiso / decisión', 'Responsable', 'Fecha límite', 'Estado']
        for indice, texto in enumerate(encabezados):
            celda = tabla_c.rows[0].cells[indice]
            celda.text = texto
            for run in celda.paragraphs[0].runs:
                run.bold = True
                run.font.color.rgb = BLANCO
                run.font.size = Pt(9.5)
            _sombrear_celda(celda, '0B5C53')
        for compromiso in compromisos:
            fila = tabla_c.add_row()
            fila.cells[0].text = compromiso.descripcion
            fila.cells[1].text = compromiso.responsable.nombre_completo
            fila.cells[2].text = str(compromiso.fecha_limite) if compromiso.fecha_limite else '—'
            estado_texto = compromiso.get_estado_display()
            if compromiso.esta_vencido:
                estado_texto += ' — Vencido'
            fila.cells[3].text = estado_texto
    else:
        _texto_o_vacio(doc, None, mensaje_vacio='Esta revisión no generó compromisos.')

    nota = doc.add_paragraph()
    nota.paragraph_format.space_before = Pt(16)
    run = nota.add_run(
        'Documento generado automáticamente desde el SGSI de ANAS WAYUU EPSI el '
        f'{timezone.localtime().strftime("%Y-%m-%d %H:%M")}. Las cifras de apoyo corresponden al momento de '
        'generación de este documento, no al momento de finalización del acta.'
    )
    run.italic = True
    run.font.size = Pt(7.5)
    run.font.color.rgb = GRIS_CLARO

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer
