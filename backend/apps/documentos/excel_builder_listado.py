# -*- coding: utf-8 -*-
"""Genera el Excel real "Listado Maestro de Documentos SGSI" (FO institucional,
aportado por el usuario) a partir de los `Documento` actuales — una hoja "Listado"
consolidada más una hoja por tipo (Manual, Política, Procedimiento, Plan, Formato,
Guía, Matriz, Programa, Protocolo, y cualquier otro tipo que llegue a tener
documentos). Se regenera en caliente en cada descarga: no hay archivo "congelado" que
se desactualice, cualquier documento guardado en el sistema aparece aquí solo.

No se usa `insert_rows`/`delete_rows` (la lección de esta sesión sobre celdas
combinadas que no se reajustan) porque esta plantilla no tiene merges por fila, solo
el título — alcanza con limpiar y reescribir celda por celda dentro de un rango
amplio."""
import io
from copy import copy
from pathlib import Path

import openpyxl

from .models import Documento

RUTA_TEMPLATE = Path(__file__).resolve().parent / 'excel_templates' / 'Listado_Maestro_de_Documentos_SGSI.xlsx'

MAPA_TIPO_HOJA = {
    Documento.Tipo.ALCANCE: 'Alcance',
    Documento.Tipo.MANUAL: 'Manual',
    Documento.Tipo.POLITICA: 'Politica',
    Documento.Tipo.PROCEDIMIENTO: 'Procedimiento',
    Documento.Tipo.PLAN: 'Plan',
    Documento.Tipo.FORMATO: 'Formato',
    Documento.Tipo.GUIA: 'Guia',
    Documento.Tipo.MATRIZ: 'Matriz',
    Documento.Tipo.PROGRAMA: 'Programa',
    Documento.Tipo.PROTOCOLO: 'Protocolo',
    Documento.Tipo.INSTRUCTIVO: 'Instructivo',
    Documento.Tipo.REGISTRO: 'Registro',
}

FILA_INICIO = 3
MAX_FILA_LIMPIAR = 500  # margen generoso sobre el volumen real actual (~60 documentos)
COLUMNAS = range(1, 10)  # A..I: PROCESO, SUB PROCESO, TIPO DE DOC., CODIGO DE PROCESO,
                         # CONSECUTIVO, CODIGO DOCUMENTO, NOMBRE DEL DOCUMENTO, VERSION,
                         # OBSERVACIONES


def _estilos_fila(ws, fila):
    return {
        col: {
            'font': copy(ws.cell(row=fila, column=col).font),
            'border': copy(ws.cell(row=fila, column=col).border),
            'fill': copy(ws.cell(row=fila, column=col).fill),
            'alignment': copy(ws.cell(row=fila, column=col).alignment),
            'number_format': ws.cell(row=fila, column=col).number_format,
        }
        for col in COLUMNAS
    }


def _limpiar(ws):
    for fila in range(FILA_INICIO, MAX_FILA_LIMPIAR + 1):
        for col in COLUMNAS:
            ws.cell(row=fila, column=col).value = None


def _partes_codigo(codigo):
    partes = codigo.split('-')
    tipo_doc = partes[0] if len(partes) > 0 else ''
    codigo_proceso = partes[1] if len(partes) > 1 else ''
    consecutivo = partes[2] if len(partes) > 2 else ''
    try:
        codigo_proceso = int(codigo_proceso)
    except ValueError:
        pass
    return tipo_doc, codigo_proceso, consecutivo


def _escribir_filas(ws, documentos, estilos):
    fila = FILA_INICIO
    for doc in documentos:
        tipo_doc, codigo_proceso, consecutivo = _partes_codigo(doc.codigo)
        valores = [
            doc.proceso, doc.sub_proceso, tipo_doc, codigo_proceso, consecutivo,
            doc.codigo, doc.titulo, doc.version_actual, '',
        ]
        for col, valor in zip(COLUMNAS, valores):
            celda = ws.cell(row=fila, column=col, value=valor)
            estilo = estilos[col]
            celda.font = estilo['font']
            celda.border = estilo['border']
            celda.fill = estilo['fill']
            celda.alignment = estilo['alignment']
            celda.number_format = estilo['number_format']
        fila += 1


def generar_listado_maestro_xlsx():
    wb = openpyxl.load_workbook(RUTA_TEMPLATE)

    documentos = list(Documento.objects.all().order_by('codigo'))
    por_tipo = {}
    for doc in documentos:
        por_tipo.setdefault(doc.tipo, []).append(doc)

    # Asegura que exista una hoja por cada tipo con al menos un documento, clonando el
    # formato de una hoja de tipo ya existente si ese tipo todavía no tenía pestaña
    # propia en la plantilla real (ej. Instructivo/Registro, sin ejemplos todavía) —
    # se hace ANTES de escribir nada, para no clonar datos ya puestos.
    hoja_base = wb['Protocolo']
    for tipo_valor, nombre_hoja in MAPA_TIPO_HOJA.items():
        if nombre_hoja not in wb.sheetnames and por_tipo.get(tipo_valor):
            nueva = wb.copy_worksheet(hoja_base)
            nueva.title = nombre_hoja
            nueva['A1'] = nombre_hoja.upper()

    ws_listado = wb['Listado']
    estilos_listado = _estilos_fila(ws_listado, FILA_INICIO)
    _limpiar(ws_listado)
    _escribir_filas(ws_listado, documentos, estilos_listado)

    for tipo_valor, nombre_hoja in MAPA_TIPO_HOJA.items():
        if nombre_hoja not in wb.sheetnames:
            continue
        ws = wb[nombre_hoja]
        estilos = _estilos_fila(ws, FILA_INICIO)
        _limpiar(ws)
        _escribir_filas(ws, por_tipo.get(tipo_valor, []), estilos)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer
